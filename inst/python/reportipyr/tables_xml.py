import os
from typing import Optional

from docx import Document
from lxml import etree

from .docx_utils import iter_cell_paragraphs
from .magic import get_magic_pattern, parse_magic_string
from .logging import setup_logger
from .util import check_duplicates, safe_resolve

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W_TBL = f"{{{W_NS}}}tbl"
W_P = f"{{{W_NS}}}p"


def add_table_xml(docx_in: str, docx_out: str, table_dir: str) -> None:
    """Insert <w:tbl> XML fragments into a docx at `{rpfy}:name.xml` magic strings.

    Reads `<table_dir>/<name>.xml` for each `.xml` magic string in the document
    and splices the parsed `<w:tbl>` element as the next sibling of the
    magic-string paragraph. Skips insertion if a `<w:tbl>` is already present
    immediately after the magic-string paragraph (idempotent re-runs).
    """
    logger = setup_logger()
    logger.debug("Starting add_table_xml function")
    logger.debug(f"Loading document from: {docx_in}")
    logger.debug(f"Table directory: {table_dir}")

    document = Document(docx_in)
    magic_pattern = get_magic_pattern()
    inserted: list[str] = []

    # Snapshot the paragraph list because insertions mutate the body tree.
    for par in list(document.paragraphs):
        matches = magic_pattern.findall(par.text)
        if not matches:
            continue
        check_duplicates(matches, "table xml names in paragraph", logger)
        for match in matches:
            for name in _xml_names_from_match(match):
                if _insert_into_body(par, name, table_dir, logger):
                    inserted.append(name)

    for cell_par, cell, _tbl in iter_cell_paragraphs(document):
        matches = magic_pattern.findall(cell_par.text)
        if not matches:
            continue
        check_duplicates(matches, "table xml names in table cell", logger)
        for match in matches:
            for name in _xml_names_from_match(match):
                if _insert_into_cell(cell_par, cell, name, table_dir, logger):
                    inserted.append(name)

    if len(set(inserted)) != len(inserted):
        logger.warning("Duplicate table XML names found in the document")

    document.save(docx_out)
    logger.info(f"Tables saved to '{docx_out}'")
    logger.debug("Exiting add_table_xml function")


def _xml_names_from_match(match: str) -> list[str]:
    args = parse_magic_string(match)
    return [
        name
        for name in args.keys()
        if os.path.splitext(name)[1].lower() == ".xml"
    ]


def _resolve_xml_path(
    table_dir: str, name: str, logger
) -> Optional[str]:
    try:
        xml_path = safe_resolve(table_dir, name)
    except ValueError:
        logger.warning(f"Path traversal blocked for: {name}")
        return None
    if not os.path.exists(xml_path):
        logger.warning(f"Table XML file not found: {xml_path}")
        return None
    return xml_path


def _load_tbl_element(xml_path: str, logger):
    try:
        tree = etree.parse(xml_path)
    except etree.XMLSyntaxError as e:
        logger.error(f"Failed to parse XML file {xml_path}: {e}")
        return None

    root = tree.getroot()
    if root.tag != W_TBL:
        logger.error(
            f"Expected root element <w:tbl> in {xml_path}, got <{root.tag}>"
        )
        return None
    return root


def _insert_into_body(par, name, table_dir, logger) -> bool:
    xml_path = _resolve_xml_path(table_dir, name, logger)
    if xml_path is None:
        return False

    par_el = par._element
    next_sibling = par_el.getnext()
    if next_sibling is not None and next_sibling.tag == W_TBL:
        logger.info(
            f"Table already present, skipping insertion for: {name}"
        )
        return False

    tbl_el = _load_tbl_element(xml_path, logger)
    if tbl_el is None:
        return False

    par_el.addnext(tbl_el)
    logger.info(f"Inserted table XML: {name}")
    return True


def _insert_into_cell(cell_par, cell, name, table_dir, logger) -> bool:
    xml_path = _resolve_xml_path(table_dir, name, logger)
    if xml_path is None:
        return False

    par_el = cell_par._element
    next_sibling = par_el.getnext()
    if next_sibling is not None and next_sibling.tag == W_TBL:
        logger.info(
            f"Cell table already present, skipping insertion for: {name}"
        )
        return False

    tbl_el = _load_tbl_element(xml_path, logger)
    if tbl_el is None:
        return False

    par_el.addnext(tbl_el)

    # OOXML requires <w:tc> to end with <w:p>. If the splice left the cell
    # ending with <w:tbl>, append an empty paragraph.
    tc_el = cell._tc
    if len(tc_el) and tc_el[-1].tag == W_TBL:
        etree.SubElement(tc_el, W_P)

    logger.info(f"Inserted cell table XML: {name}")
    return True
