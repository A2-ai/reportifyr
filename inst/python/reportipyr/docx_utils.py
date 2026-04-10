from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .magic import get_magic_pattern

CAPTION_STYLE = "Caption"


def iter_cell_paragraphs(doc):
    """Yield (paragraph, cell, table_element) for paragraphs inside table cells.

    Deduplicates merged cells using cell._tc identity. Stores references
    to _tc elements (not just id()) to prevent lxml proxy garbage collection
    from causing id reuse across iterations.
    """
    for table in doc.tables:
        seen_tcs = {}  # id -> _tc reference (prevents GC/id reuse)
        for row in table.rows:
            for cell in row.cells:
                tc = cell._tc
                tc_id = id(tc)
                if tc_id in seen_tcs:
                    continue
                seen_tcs[tc_id] = tc
                for para in cell.paragraphs:
                    yield para, cell, table._element


def keep_caption_next(docx_in, docx_out):
    doc = Document(docx_in)
    paras = doc.paragraphs
    n = len(paras)

    magic_pattern = get_magic_pattern()

    for i, p in enumerate(paras):
        is_caption = (p.style and p.style.name == CAPTION_STYLE) or any(
            ("SEQ Table" in instr.text or "SEQ Figure" in instr.text)
            for instr in p._element.xpath(".//w:instrText")
        )
        if not is_caption:
            continue

        # add keepNext to caption
        pPr = p._element.get_or_add_pPr()
        if not pPr.xpath("./w:keepNext"):
            pPr.append(OxmlElement("w:keepNext"))

        # scan forward for first paragraph with magic string
        for j in range(i + 1, n):
            q = paras[j]
            has_magic = bool(
                magic_pattern.search(
                    "".join(t.text for t in q._element.xpath(".//w:t"))
                )
            )
            if has_magic:
                qPr = q._element.get_or_add_pPr()
                if not qPr.xpath("./w:keepNext"):
                    qPr.append(OxmlElement("w:keepNext"))
                break

    doc.save(docx_out)


def remove_bookmarks(docx_in, docx_out):
    doc = Document(docx_in)

    fp_bookmark_ids = set()

    for element in doc.element.findall(
        ".//w:bookmarkStart", namespaces=doc.element.nsmap
    ):
        bookmark_name = element.get(qn("w:name"))
        if bookmark_name and bookmark_name.startswith("fp_"):
            fp_bookmark_ids.add(element.get(qn("w:id")))
            parent = element.getparent()
            parent.remove(element)

    for element in doc.element.findall(
        ".//w:bookmarkEnd", namespaces=doc.element.nsmap
    ):
        bookmark_id = element.get(qn("w:id"))
        if bookmark_id in fp_bookmark_ids:
            parent = element.getparent()
            parent.remove(element)

    doc.save(docx_out)
