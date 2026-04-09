import os
import tempfile
from typing import Optional

from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from PIL import Image, ImageDraw, ImageFont

from .alt_text import is_artifact_unchanged
from .config import load_yaml
from .magic import get_magic_pattern, parse_magic_string
from .logging import setup_logger
from .util import check_duplicates, create_label, safe_resolve


def add_figure(
    docx_in: str,
    docx_out: str,
    figure_dir: str,
    config_yaml: Optional[str],
    fig_width: Optional[float] = None,
    fig_height: Optional[float] = None,
):
    logger = setup_logger()
    logger.debug("Starting add_figure function")
    logger.debug(f"Loading document from: {docx_in}")
    logger.debug(f"Figure directory: {figure_dir}")

    document = Document(docx_in)

    # load config.yaml or set empty dict for defaults.
    if config_yaml is not None:
        logger.debug(f"Loading config from: {config_yaml}")
        config = load_yaml(config_yaml)
    else:
        config = {}

    magic_pattern = get_magic_pattern()

    found_magic_strings = []

    # Process paragraphs in reverse order to avoid index shifting issues
    for i, par in enumerate(reversed(document.paragraphs)):
        # Calculate the actual index in the original document
        actual_index = len(document.paragraphs) - 1 - i

        matches = magic_pattern.findall(par.text)
        if matches:
            check_duplicates(matches, f"figure names in paragraph {actual_index+1}", logger)

            # Check if figures already exist (skip_unchanged kept them)
            # Count consecutive drawing paragraphs after the magic string
            # in the live XML tree and compare to expected PNG count.
            parent = par._element.getparent()
            body_children = list(parent)
            magic_idx = body_children.index(par._element)
            existing_drawings = 0
            for offset in range(1, len(body_children) - magic_idx):
                sibling = body_children[magic_idx + offset]
                if (
                    sibling.tag.endswith("}p")
                    and not "".join(
                        t.text for t in sibling.xpath(".//w:t") if t.text
                    ).strip()
                    and sibling.xpath(".//w:drawing")
                ):
                    existing_drawings += 1
                else:
                    break

            all_figure_args = parse_magic_string(matches[0])
            png_count = sum(
                1 for f in all_figure_args
                if os.path.splitext(f)[1].lower() == ".png"
            )
            if existing_drawings >= png_count > 0:
                logger.info(
                    f"Figures already present for: "
                    f"{list(all_figure_args.keys())}, skipping"
                )
                for figure in all_figure_args:
                    found_magic_strings.append(figure)
                continue

            for match in matches:
                logger.debug(f"Processing magic string: {match}")
                # Extract the image directory from the match
                # figure_name now contains potentially a list
                # of file names and args. like
                # [file.ext, file2.ext]<width: 4, height: 6>
                figure_args = parse_magic_string(match)

                if len(figure_args) > 1:
                    figures = list(reversed(figure_args.keys()))
                else:
                    figures = list(figure_args.keys())

                for fig_idx, figure in enumerate(figures):
                    extension = os.path.splitext(figure)[1].lower()
                    if extension != ".png":
                        logger.debug(f"Skipping non-png file: {figure}")
                        continue

                    logger.info(f"Processing figure: {figure}")

                    add_label = False
                    if len(figures) > 1 and config.get("label_multi_figures", False):
                        add_label = True

                    found_magic_strings.append(figure)
                    try:
                        image_path = safe_resolve(figure_dir, figure)
                    except ValueError:
                        logger.warning(f"Path traversal blocked for: {figure}")
                        continue
                    if os.path.exists(image_path):
                        if add_label:
                            # since list is reversed need to use correct
                            # index, index 0 corresponds to the last
                            # element in list so len(figures) - 1 for 1-index
                            labeled_image = add_label_to_image(
                                image_path, len(figures) - fig_idx - 1, logger
                            )
                        else:
                            labeled_image = image_path

                        # Insert new paragraph after the current paragraph
                        new_par = document.add_paragraph()
                        run = new_par.add_run()

                        # Move paragraph to correct position (after current paragraph)
                        parent = par._element.getparent()
                        new_par._element.getparent().remove(new_par._element)
                        target_index = list(parent).index(par._element)
                        parent.insert(target_index + 1, new_par._element)

                        # Configure image size
                        if config.get("use_embedded_dimensions", True) and set(
                            figure_args[figure].keys()
                        ).intersection(["width", "height"]):
                            embedded_width = figure_args[figure].get("width")
                            embedded_height = figure_args[figure].get("height")
                            logger.debug(
                                f"Using embedded size (width={embedded_width}, height={embedded_height})"
                            )
                            run.add_picture(
                                labeled_image,
                                width=(
                                    Inches(float(embedded_width))
                                    if embedded_width
                                    else None
                                ),
                                height=(
                                    Inches(float(embedded_height))
                                    if embedded_height
                                    else None
                                ),
                            )

                        elif config.get("use_artifact_size", True):
                            logger.debug(
                                "Using artifact size (original image dimensions)"
                            )
                            run.add_picture(labeled_image)

                        else:
                            default_width = config.get("default_fig_width", 6)
                            if set(figure_args[figure].keys()).intersection(
                                ["width", "height"]
                            ):
                                logger.debug("Using embedded size from magic string")
                                embedded_width = figure_args[figure].get("width")
                                embedded_height = figure_args[figure].get("height")
                                run.add_picture(
                                    labeled_image,
                                    width=(
                                        Inches(float(embedded_width))
                                        if embedded_width
                                        else None
                                    ),
                                    height=(
                                        Inches(float(embedded_height))
                                        if embedded_height
                                        else None
                                    ),
                                )

                            elif fig_width is not None and fig_height is not None:
                                logger.debug(
                                    f"Using CLI arguments (width={fig_width}, height={fig_height})"
                                )
                                run.add_picture(
                                    labeled_image,
                                    width=Inches(fig_width),
                                    height=Inches(fig_height),
                                )
                            elif fig_width is not None:
                                logger.debug(f"Using CLI width argument: {fig_width}")
                                run.add_picture(labeled_image, width=Inches(fig_width))
                            elif fig_height is not None:
                                logger.debug(f"Using CLI height argument: {fig_height}")
                                run.add_picture(
                                    labeled_image, height=Inches(fig_height)
                                )
                            else:
                                logger.debug(f"Using default width: {default_width}")
                                run.add_picture(
                                    labeled_image, width=Inches(default_width)
                                )

                        # Set alignment
                        alignment = config.get("fig_alignment", "center").lower()
                        match alignment:
                            case "center":
                                new_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            case "left":
                                new_par.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            case "right":
                                new_par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                            case _:
                                new_par.alignment = WD_ALIGN_PARAGRAPH.CENTER

                        logger.info(f"Inserted figure: {figure}")
                    else:
                        logger.warning(f"Figure file not found: {image_path}")

    if len(set(found_magic_strings)) != len(found_magic_strings):
        logger.warning("Duplicate figure names found in the document")

    document.save(docx_out)
    logger.info(f"Figures saved to '{docx_out}'")
    logger.debug("Exiting add_figure function")


def add_label_to_image(image_path: str, index: int, logger) -> str:
    """
    This function takes in a path to an image and an index
    and adds the corresponding letter to the image upper
    left corner. If the index is > 25 then the label will
    use multiple letters.
    26 = AA
    27 = AB, etc.

    This function saves the updated image to tmp and returns
    the path to the temp image.
    """
    label = create_label(index)

    # load in image and create draw object
    # and set font
    try:
        img = Image.open(image_path)
    except Exception as e:
        logger.error(f"Failed to open image {image_path}: {e}")
        raise

    draw = ImageDraw.Draw(img)
    font = ImageFont.load_default(size=56)

    img_width, img_height = img.size
    original_format = img.format or os.path.splitext(image_path)[1][1:].upper()
    original_dpi = img.info.get("dpi", (72, 72))  # word default
    logger.debug(
        f"Image dimensions: {img_width}x{img_height}, format: {original_format}"
    )

    left_corner_position = (20, 20)

    draw.rectangle(
        xy=[
            left_corner_position[0] - 5,
            left_corner_position[1] - 5,
            left_corner_position[0] + img_width * 0.025,
            left_corner_position[1] + img_height * 0.025,
        ],
        fill=(255, 255, 255),
    )

    draw.text(left_corner_position, label, fill=(0, 0, 0), font=font)
    temp_file = tempfile.NamedTemporaryFile(
        delete=False, suffix=os.path.splitext(image_path)[1]
    )
    temp_path = temp_file.name
    temp_file.close()

    try:
        img.save(temp_path, format=original_format, dpi=original_dpi)
    except Exception as e:
        logger.error(f"Failed to save labeled image: {e}")
        raise

    return temp_path


def remove_figures(
    docx_in: str,
    docx_out: str,
    config_yaml: Optional[str],
    figure_dir: str | None = None,
):
    logger = setup_logger()
    doc = Document(docx_in)
    paragraphs = doc.paragraphs

    # load config.yaml or set empty dict for defaults.
    if config_yaml is not None:
        config = load_yaml(config_yaml)
    else:
        config = {}

    skip_unchanged = config.get("skip_unchanged", False)

    for i, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if text.startswith("{rpfy}:"):
            figure_args = parse_magic_string(text)

            # Check alt text hashes for skip_unchanged
            if skip_unchanged and figure_dir:
                # Check next paragraphs for drawings with alt text
                all_unchanged = True
                for j, fig_name in enumerate(figure_args.keys()):
                    if i + j + 1 < len(paragraphs):
                        next_par = paragraphs[i + j + 1]
                        drawings = next_par._element.xpath(
                            ".//w:drawing"
                        )
                        if drawings:
                            for d in drawings:
                                for inline in d.xpath(
                                    ".//wp:inline"
                                ):
                                    for dp in inline.xpath(
                                        ".//wp:docPr"
                                    ):
                                        alt = dp.get("descr", "")
                                        if not is_artifact_unchanged(
                                            alt, figure_dir, fig_name
                                        ):
                                            all_unchanged = False
                        else:
                            all_unchanged = False
                    else:
                        all_unchanged = False

                if all_unchanged:
                    logger.info(
                        f"Skipping removal of unchanged figures: "
                        f"{list(figure_args.keys())}"
                    )
                    continue

            update_magic_string = False

            paragraphs_to_remove = []
            for j, args in enumerate(figure_args.values()):
                if i + j + 1 < len(paragraphs):
                    next_par = paragraphs[i + j + 1]
                    if not next_par.text.strip() and next_par._element.xpath(
                        ".//w:drawing"
                    ):
                        paragraphs_to_remove.append((i + j + 1, next_par))
                        if config.get("use_embedded_dimensions", True):
                            dimensions = get_figure_dimensions(next_par)
                            # set width and height from emu to Inches
                            if dimensions.get("width"):
                                args["width"] = str(
                                    round(dimensions["width"] / 914400, 2)
                                )
                                update_magic_string = True
                            if dimensions.get("height"):
                                args["height"] = str(
                                    round(dimensions["height"] / 914400, 2)
                                )
                                update_magic_string = True

            if update_magic_string:
                new_magic_string = "{rpfy}:"
                if len(figure_args) > 1:
                    new_magic_string += "["
                    ending_string = "]"
                else:
                    ending_string = ""
                for fig_idx, (fig, arg) in enumerate(figure_args.items()):
                    arg_string = "<"
                    for p_idx, (prop, val) in enumerate(arg.items()):
                        arg_string += f"{prop}: {val}"
                        if p_idx + 1 != len(arg):
                            arg_string += ", "
                    arg_string += ">"

                    new_magic_string += f"{fig}{arg_string}"
                    if fig_idx + 1 != len(figure_args):
                        new_magic_string += ", "

                new_magic_string += ending_string
                paragraph.text = new_magic_string

            for _, par in reversed(paragraphs_to_remove):
                par._element.getparent().remove(par._element)

    doc.save(docx_out)


def get_figure_dimensions(paragraph) -> dict[str, Optional[int]]:
    """Extract width and height from a paragraph containing a drawing."""
    # Get all drawing elements in the paragraph
    drawing_elements = paragraph._element.xpath(".//w:drawing")

    if not drawing_elements:
        return {}

    # Get the first drawing element
    drawing = drawing_elements[0]

    # Find the extent element within the drawing
    extent_elements = drawing.xpath(".//wp:extent")

    if not extent_elements:
        return {}

    # Get the first extent element
    extent = extent_elements[0]

    # Extract cx and cy values
    cx = extent.get(
        "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}cx"
    )
    cy = extent.get(
        "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}cy"
    )

    # If using namespaces doesn't work, try getting attributes directly
    if cx is None:
        cx = extent.get("cx")
    if cy is None:
        cy = extent.get("cy")

    if cx is None or cy is None:
        return {}

    # Convert to integers
    width = int(cx)
    height = int(cy)

    return {"width": width, "height": height}


def add_path_overlay_to_image(image_path: str, source_text: str) -> None:
    """Stamp a source-path label onto the bottom-left of an image, in-place."""
    logger = setup_logger()
    logger.debug(f"Adding path overlay to {image_path}: {source_text}")

    if not image_path.lower().endswith(".png"):
        logger.warning(
            f"Skipping path overlay for non-png format: {image_path}"
        )
        return

    img = Image.open(image_path)
    draw = ImageDraw.Draw(img)
    original_format = (
        img.format or os.path.splitext(image_path)[1][1:].upper()
    )
    original_dpi = img.info.get("dpi", (72, 72))
    img_width, img_height = img.size

    font_size = max(12, int(min(img_width, img_height) * 0.02))
    font = ImageFont.load_default(size=font_size)

    padding = 10
    text_position = (padding, img_height - padding - font_size)
    draw.text(
        text_position,
        f"source: {source_text}",
        fill=(0, 0, 0),
        font=font,
    )

    img.save(image_path, format=original_format, dpi=original_dpi)
    logger.info(f"Path overlay saved to {image_path}")
