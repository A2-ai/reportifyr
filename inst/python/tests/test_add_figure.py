import tempfile
from pathlib import Path

from docx import Document
from PIL import Image

from reportipyr.figures import add_figure


def _make_png(path: Path) -> None:
    img = Image.new("RGB", (10, 10), color=(0, 0, 255))
    img.save(path)


def _get_first_extent(doc: Document) -> tuple[int, int] | None:
    for p in doc.paragraphs:
        extents = p._element.xpath(".//*[local-name()='extent']")
        if extents:
            extent = extents[0]
            cx = extent.get(
                "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}cx"
            ) or extent.get("cx")
            cy = extent.get(
                "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}cy"
            ) or extent.get("cy")
            if cx is not None and cy is not None:
                return int(cx), int(cy)
    return None


def test_add_figure_respects_cli_dimensions():
    figure_dir = Path(tempfile.mkdtemp())
    img_path = figure_dir / "figure.png"
    _make_png(img_path)

    # use_artifact_size defaults to True, so we must disable it
    # to reach the CLI dimension branch
    config_yaml = _write_yaml("use_artifact_size: false")

    docx_in = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    doc = Document()
    doc.add_paragraph("{rpfy}:figure.png")
    doc.save(docx_in)

    docx_out = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    add_figure(
        docx_in=docx_in,
        docx_out=docx_out,
        figure_dir=str(figure_dir),
        config_yaml=config_yaml,
        fig_width=4.0,
        fig_height=5.0,
    )

    out_doc = Document(docx_out)
    extent = _get_first_extent(out_doc)
    assert extent is not None

    expected_width = int(4.0 * 914400)
    expected_height = int(5.0 * 914400)
    assert extent[0] == expected_width
    assert extent[1] == expected_height


def _write_yaml(contents: str) -> str:
    path = tempfile.NamedTemporaryFile(suffix=".yaml", delete=False).name
    Path(path).write_text(contents)
    return path


def test_add_figure_respects_embedded_size():
    figure_dir = Path(tempfile.mkdtemp())
    img_path = figure_dir / "figure.png"
    _make_png(img_path)

    docx_in = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    doc = Document()
    doc.add_paragraph("{rpfy}:figure.png<width: 4, height: 5>")
    doc.save(docx_in)

    config_yaml = _write_yaml("use_embedded_dimensions: true\n")

    docx_out = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    add_figure(
        docx_in=docx_in,
        docx_out=docx_out,
        figure_dir=str(figure_dir),
        config_yaml=config_yaml,
    )

    out_doc = Document(docx_out)
    extent = _get_first_extent(out_doc)
    assert extent is not None

    expected_width = int(4.0 * 914400)
    expected_height = int(5.0 * 914400)
    assert extent[0] == expected_width
    assert extent[1] == expected_height
