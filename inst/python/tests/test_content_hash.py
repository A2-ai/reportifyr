"""
Programmatic test suite for content hash validation.

Replicates the empirical test cases documented in hash_ordering_results.md.
Each test constructs before/after table XML via python-docx + lxml,
then asserts hash equality/inequality using compute_table_content_hash().

Baseline table:
    C1       C2       C3
    Metric   Quarter  Value
    Revenue  Q1       100
    Revenue  Q2       120
"""

from copy import deepcopy

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportipyr.alt_text import (
    compute_table_content_hash,
    _decode_body_grid,
    _detect_header_rows,
    encode_body_grid,
    extract_body_grid,
    grid_to_canonical,
    _is_full_width_row,
    _logical_column_count,
)
from reportipyr.tables import _update_cell_text, reconcile_table_cells


# ── Helpers ────────────────────────────────────────────────────────────────


BASELINE_DATA = [
    ["C1", "C2", "C3"],
    ["Metric", "Quarter", "Value"],
    ["Revenue", "Q1", "100"],
    ["Revenue", "Q2", "120"],
]


def _make_baseline_table():
    """Create a Document with the baseline 4x3 table and return (doc, tbl_element)."""
    doc = Document()
    table = doc.add_table(rows=4, cols=3)
    for r, row_data in enumerate(BASELINE_DATA):
        for c, val in enumerate(row_data):
            table.cell(r, c).paragraphs[0].text = val
    return doc, table._tbl


def _baseline_hash():
    """Return the hash of an unmodified baseline table."""
    _, tbl = _make_baseline_table()
    return compute_table_content_hash(tbl)


def _get_cell(tbl, row, col):
    """Return the w:tc element at (row, col) — 0-indexed."""
    tr = tbl.findall(qn("w:tr"))[row]
    return tr.findall(qn("w:tc"))[col]


def _get_wt_elements(tc):
    """Return all w:t elements within a cell."""
    return list(tc.iter(qn("w:t")))


def _get_run_elements(tc):
    """Return all w:r elements within a cell."""
    return list(tc.iter(qn("w:r")))


def _get_paragraph_elements(tc):
    """Return all w:p elements that are direct children of the cell."""
    return tc.findall(qn("w:p"))


# ── Hash should NOT change ─────────────────────────────────────────────────


def test_e01_split_run():
    """E01: Split run should not change value-based hash."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    # Mutate: split Revenue run into two runs "Rev" + "enue"
    _, tbl2 = _make_baseline_table()
    tc = _get_cell(tbl2, 2, 0)  # row 3, col 1
    p = _get_paragraph_elements(tc)[0]

    # Remove existing run(s)
    for r in list(p.findall(qn("w:r"))):
        p.remove(r)

    # Add two new runs
    for text in ["Rev", "enue"]:
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        run.append(t)
        p.append(run)

    assert compute_table_content_hash(tbl2) == base


def test_e01_followup_split_run_no_bold():
    """E01 follow-up: Split run without bold should not change hash."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    tc = _get_cell(tbl2, 2, 0)
    p = _get_paragraph_elements(tc)[0]

    for r in list(p.findall(qn("w:r"))):
        p.remove(r)

    # Two runs, explicitly no bold
    for text in ["Rev", "enue"]:
        run = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        b = OxmlElement("w:b")
        b.set(qn("w:val"), "0")
        rPr.append(b)
        run.append(rPr)
        t = OxmlElement("w:t")
        t.text = text
        run.append(t)
        p.append(run)

    assert compute_table_content_hash(tbl2) == base


def test_e02_merge_runs():
    """E02: Identical single-run cell in both — identity check."""
    _, tbl1 = _make_baseline_table()
    _, tbl2 = _make_baseline_table()
    assert compute_table_content_hash(tbl1) == compute_table_content_hash(tbl2)


def test_e03_two_paragraphs_in_cell():
    """E03: Two paragraphs in a cell (same visible text) should not change hash."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    tc = _get_cell(tbl2, 2, 0)

    # Remove existing paragraph
    for p in list(_get_paragraph_elements(tc)):
        tc.remove(p)

    # Add two paragraphs: "Rev" and "enue"
    for text in ["Rev", "enue"]:
        p = OxmlElement("w:p")
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        run.append(t)
        p.append(run)
        tc.append(p)

    assert compute_table_content_hash(tbl2) == base


def test_e05_header_semantics():
    """E05: Adding w:tblHeader excludes that row from hash (hash changes).

    Under the new body-only hashing contract, marking a row as a header
    causes it to be skipped. This is intentional — the hash now reflects
    only body data rows.
    """
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    # Add w:tblHeader to row 1
    tr = tbl2.findall(qn("w:tr"))[0]
    trPr = tr.find(qn("w:trPr"))
    if trPr is None:
        trPr = OxmlElement("w:trPr")
        tr.insert(0, trPr)
    header = OxmlElement("w:tblHeader")
    trPr.append(header)

    assert compute_table_content_hash(tbl2) != base


def test_e06_style_only():
    """E06: Style-only changes (bold/italic/underline) should not change hash."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()

    # Bold a run in row 3, col 1
    tc = _get_cell(tbl2, 2, 0)
    runs = _get_run_elements(tc)
    if runs:
        rPr = OxmlElement("w:rPr")
        rPr.append(OxmlElement("w:b"))
        runs[0].insert(0, rPr)

    # Italic a run in row 3, col 2
    tc2 = _get_cell(tbl2, 2, 1)
    runs2 = _get_run_elements(tc2)
    if runs2:
        rPr2 = OxmlElement("w:rPr")
        rPr2.append(OxmlElement("w:i"))
        runs2[0].insert(0, rPr2)

    # Underline a run in row 3, col 3
    tc3 = _get_cell(tbl2, 2, 2)
    runs3 = _get_run_elements(tc3)
    if runs3:
        rPr3 = OxmlElement("w:rPr")
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr3.append(u)
        runs3[0].insert(0, rPr3)

    assert compute_table_content_hash(tbl2) == base


def test_e06_5_identical_baseline():
    """E6.5: Two identical baseline tables should have equal hashes."""
    _, tbl1 = _make_baseline_table()
    _, tbl2 = _make_baseline_table()
    assert compute_table_content_hash(tbl1) == compute_table_content_hash(tbl2)


def test_e13_trailing_space():
    """E13: Trailing space should not change hash (cell-level strip)."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    tc = _get_cell(tbl2, 2, 2)  # row 3, col 3 = "100"
    wt = _get_wt_elements(tc)[0]
    wt.text = "100 "
    wt.set(qn("xml:space"), "preserve")

    assert compute_table_content_hash(tbl2) == base


def test_e16_5_tab_element():
    """E16.5: w:tab insertion should not change hash."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    tc = _get_cell(tbl2, 2, 0)  # row 3, col 1
    runs = _get_run_elements(tc)
    if runs:
        # Insert w:tab before the w:t element
        tab = OxmlElement("w:tab")
        runs[0].insert(0, tab)

    assert compute_table_content_hash(tbl2) == base


def test_e20_trailing_nbsp():
    """E20: Trailing non-breaking space should not change hash (.strip() removes it)."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    tc = _get_cell(tbl2, 2, 2)  # row 3, col 3 = "100"
    wt = _get_wt_elements(tc)[0]
    wt.text = "100\u00a0"

    assert compute_table_content_hash(tbl2) == base


def test_e23_line_break():
    """E23: w:br insertion should not change hash."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    tc = _get_cell(tbl2, 2, 0)  # row 3, col 1
    runs = _get_run_elements(tc)
    if runs:
        br = OxmlElement("w:br")
        runs[0].append(br)

    assert compute_table_content_hash(tbl2) == base


# ── Hash SHOULD change ────────────────────────────────────────────────────


def test_e04_value_change():
    """E04: Actual value change ('Revenue' -> 'Rev enue') should change hash."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    tc = _get_cell(tbl2, 2, 0)
    wt = _get_wt_elements(tc)[0]
    wt.text = "Rev enue"

    assert compute_table_content_hash(tbl2) != base


def test_e07_horizontal_merge():
    """E07: Horizontal cell merge collapses value sequence — hash should change.

    Documents known behavior: merging C1+C2 produces 'C1C2' in one cell.
    """
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    tr = tbl2.findall(qn("w:tr"))[0]  # row 1
    tcs = tr.findall(qn("w:tc"))

    # Merge cell 0 and cell 1: set gridSpan=2 on cell 0, append cell 1 text, remove cell 1
    tc0 = tcs[0]
    tc1 = tcs[1]

    # Get text from cell 1
    tc1_text = "".join(t.text for t in tc1.iter(qn("w:t")) if t.text)

    # Append cell 1 text to cell 0
    wt0 = list(tc0.iter(qn("w:t")))[0]
    wt0.text = wt0.text + tc1_text

    # Set gridSpan
    tcPr = tc0.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = OxmlElement("w:tcPr")
        tc0.insert(0, tcPr)
    gridSpan = OxmlElement("w:gridSpan")
    gridSpan.set(qn("w:val"), "2")
    tcPr.append(gridSpan)

    # Remove cell 1
    tr.remove(tc1)

    assert compute_table_content_hash(tbl2) != base


def test_e08_spanner_label():
    """E08: Merged spanner with new label text should change hash."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    tr = tbl2.findall(qn("w:tr"))[0]
    tcs = tr.findall(qn("w:tc"))

    tc0 = tcs[0]
    tc1 = tcs[1]

    # Set cell 0 text to "FY2026"
    wt0 = list(tc0.iter(qn("w:t")))[0]
    wt0.text = "FY2026"

    # Set gridSpan=2
    tcPr = tc0.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = OxmlElement("w:tcPr")
        tc0.insert(0, tcPr)
    gridSpan = OxmlElement("w:gridSpan")
    gridSpan.set(qn("w:val"), "2")
    tcPr.append(gridSpan)

    # Remove cell 1
    tr.remove(tc1)

    assert compute_table_content_hash(tbl2) != base


def test_e09_row_swap():
    """E09: Swapping data row order should change hash."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    trs = tbl2.findall(qn("w:tr"))

    # Swap rows 2 and 3 (0-indexed: rows at index 2 and 3)
    row3 = trs[2]
    row4 = trs[3]
    tbl2.remove(row4)
    tbl2.insert(list(tbl2).index(row3), row4)

    assert compute_table_content_hash(tbl2) != base


def test_e10_column_swap():
    """E10: Swapping column order should change hash."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    # Swap columns 0 and 1 in every row
    for tr in tbl2.findall(qn("w:tr")):
        tcs = tr.findall(qn("w:tc"))
        tc0 = tcs[0]
        tc1 = tcs[1]
        tr.remove(tc1)
        tr.insert(list(tr).index(tc0), tc1)

    assert compute_table_content_hash(tbl2) != base


def test_e11_numeric_change():
    """E11: Numeric text change ('100' -> '100.0') should change hash."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    tc = _get_cell(tbl2, 2, 2)  # row 3, col 3
    wt = _get_wt_elements(tc)[0]
    wt.text = "100.0"

    assert compute_table_content_hash(tbl2) != base


def test_e12_numeric_format():
    """E12: Numeric formatting variant ('100' -> '1,00') should change hash."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    tc = _get_cell(tbl2, 2, 2)
    wt = _get_wt_elements(tc)[0]
    wt.text = "1,00"

    assert compute_table_content_hash(tbl2) != base


def test_e15_empty_to_na():
    """E15: Empty cell to 'NA' should change hash."""
    # Before: C3 header is empty
    doc1 = Document()
    table1 = doc1.add_table(rows=4, cols=3)
    data_with_empty = [
        ["C1", "C2", ""],
        ["Metric", "Quarter", "Value"],
        ["Revenue", "Q1", "100"],
        ["Revenue", "Q2", "120"],
    ]
    for r, row_data in enumerate(data_with_empty):
        for c, val in enumerate(row_data):
            table1.cell(r, c).paragraphs[0].text = val
    hash_empty = compute_table_content_hash(table1._tbl)

    # After: C3 header is "NA"
    doc2 = Document()
    table2 = doc2.add_table(rows=4, cols=3)
    data_with_na = [
        ["C1", "C2", "NA"],
        ["Metric", "Quarter", "Value"],
        ["Revenue", "Q1", "100"],
        ["Revenue", "Q2", "120"],
    ]
    for r, row_data in enumerate(data_with_na):
        for c, val in enumerate(row_data):
            table2.cell(r, c).paragraphs[0].text = val
    hash_na = compute_table_content_hash(table2._tbl)

    assert hash_empty != hash_na


def test_e16_add_empty_row():
    """E16: Adding an empty row should change hash."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    # Prepend an empty 3-cell row
    new_tr = OxmlElement("w:tr")
    for _ in range(3):
        tc = OxmlElement("w:tc")
        p = OxmlElement("w:p")
        tc.append(p)
        new_tr.append(tc)

    # Insert at beginning (before first row)
    first_tr = tbl2.findall(qn("w:tr"))[0]
    tbl2.insert(list(tbl2).index(first_tr), new_tr)

    assert compute_table_content_hash(tbl2) != base


def test_e19_vertical_merge():
    """E19: Vertical merge alters effective values — hash should change.

    After merge: row 3 col 1 = 'Revenue' (restart), row 4 col 1 = '' (continue).
    """
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()

    # Row 3 (index 2), col 1 (index 0): set vMerge restart
    tc_restart = _get_cell(tbl2, 2, 0)
    tcPr = tc_restart.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = OxmlElement("w:tcPr")
        tc_restart.insert(0, tcPr)
    vm = OxmlElement("w:vMerge")
    vm.set(qn("w:val"), "restart")
    tcPr.append(vm)

    # Row 4 (index 3), col 1 (index 0): set vMerge continue and clear text
    tc_cont = _get_cell(tbl2, 3, 0)
    tcPr2 = tc_cont.find(qn("w:tcPr"))
    if tcPr2 is None:
        tcPr2 = OxmlElement("w:tcPr")
        tc_cont.insert(0, tcPr2)
    vm2 = OxmlElement("w:vMerge")
    tcPr2.append(vm2)

    # Clear text in continuation cell
    for wt in _get_wt_elements(tc_cont):
        wt.text = ""

    assert compute_table_content_hash(tbl2) != base


def test_e21_em_to_en_dash():
    """E21: Em dash to en dash substitution should change hash."""
    # Before: Q—1 (em dash U+2014)
    doc1 = Document()
    table1 = doc1.add_table(rows=4, cols=3)
    data_em = deepcopy(BASELINE_DATA)
    data_em[2][1] = "Q\u20141"
    for r, row_data in enumerate(data_em):
        for c, val in enumerate(row_data):
            table1.cell(r, c).paragraphs[0].text = val
    hash_em = compute_table_content_hash(table1._tbl)

    # After: Q–1 (en dash U+2013)
    doc2 = Document()
    table2 = doc2.add_table(rows=4, cols=3)
    data_en = deepcopy(BASELINE_DATA)
    data_en[2][1] = "Q\u20131"
    for r, row_data in enumerate(data_en):
        for c, val in enumerate(row_data):
            table2.cell(r, c).paragraphs[0].text = val
    hash_en = compute_table_content_hash(table2._tbl)

    assert hash_em != hash_en


def test_e22_smart_quotes():
    """E22: Smart/curly quotes vs straight quotes should change hash."""
    # Before: smart quotes
    doc1 = Document()
    table1 = doc1.add_table(rows=4, cols=3)
    data_smart = deepcopy(BASELINE_DATA)
    data_smart[1][2] = "\u201cValue\u201d"
    for r, row_data in enumerate(data_smart):
        for c, val in enumerate(row_data):
            table1.cell(r, c).paragraphs[0].text = val
    hash_smart = compute_table_content_hash(table1._tbl)

    # After: straight quotes
    doc2 = Document()
    table2 = doc2.add_table(rows=4, cols=3)
    data_straight = deepcopy(BASELINE_DATA)
    data_straight[1][2] = '"Value"'
    for r, row_data in enumerate(data_straight):
        for c, val in enumerate(row_data):
            table2.cell(r, c).paragraphs[0].text = val
    hash_straight = compute_table_content_hash(table2._tbl)

    assert hash_smart != hash_straight


def test_e24_zwsp():
    """E24: Zero-width space insertion should change hash."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    tc = _get_cell(tbl2, 2, 1)  # row 3, col 2 = "Q1"
    wt = _get_wt_elements(tc)[0]
    wt.text = "Q1\u200b"

    assert compute_table_content_hash(tbl2) != base


# ── Header / Footer Skip Tests ────────────────────────────────────────────

# Helper to build a pharma-style param table:
#   Row 0: "" (span=3) | "Estimate" | "95% CI" | "Shrinkage (%)"  [header]
#   Row 1: "Structural Model Parameters" (span=6)                 [full-width]
#   Row 2: "CL/F" | "θ1" | "Clearance" | "170" | "105, 235" | "--" [body]
#   Row 3: "Vc/F" | "θ2" | "Volume"    | "40"  | "34, 46"   | "--" [body]
#   Row 4: "Abbreviations: CI = ..." (span=6)                     [full-width]

PARAM_HEADER = [
    (["", "Estimate", "95% CI", "Shrinkage (%)"], [(3, 0), (1, 1), (1, 2), (1, 3)]),
]
PARAM_SECTION = [
    (["Structural Model Parameters"], [(6, 0)]),
]
PARAM_BODY = [
    (["CL/F", "\u03b81", "Clearance", "170", "105, 235", "--"], None),
    (["Vc/F", "\u03b82", "Volume", "40", "34, 46", "--"], None),
]
PARAM_FOOTER = [
    (["Abbreviations: CI = confidence intervals"], [(6, 0)]),
]


def _make_param_table(with_tbl_header=False):
    """Build a 6-column param table with header spans, section label, data, footer."""
    doc = Document()
    table = doc.add_table(rows=0, cols=6)
    tbl = table._tbl

    all_rows = PARAM_HEADER + PARAM_SECTION + PARAM_BODY + PARAM_FOOTER

    for row_idx, (texts, spans) in enumerate(all_rows):
        tr = OxmlElement("w:tr")

        # Set w:tblHeader on header rows if requested
        if with_tbl_header and row_idx < len(PARAM_HEADER):
            trPr = OxmlElement("w:trPr")
            trPr.append(OxmlElement("w:tblHeader"))
            tr.append(trPr)

        if spans:
            # Row with gridSpan cells
            for text, (span, _) in zip(texts, spans):
                tc = OxmlElement("w:tc")
                if span > 1:
                    tcPr = OxmlElement("w:tcPr")
                    gs = OxmlElement("w:gridSpan")
                    gs.set(qn("w:val"), str(span))
                    tcPr.append(gs)
                    tc.append(tcPr)
                p = OxmlElement("w:p")
                run = OxmlElement("w:r")
                t = OxmlElement("w:t")
                t.text = text
                run.append(t)
                p.append(run)
                tc.append(p)
                tr.append(tc)
        else:
            # Normal row — one cell per column
            for text in texts:
                tc = OxmlElement("w:tc")
                p = OxmlElement("w:p")
                run = OxmlElement("w:r")
                t = OxmlElement("w:t")
                t.text = text
                run.append(t)
                p.append(run)
                tc.append(p)
                tr.append(tc)

        tbl.append(tr)

    return doc, tbl


def test_header_skip_with_tbl_header():
    """Header rows with w:tblHeader are excluded from hash."""
    _, tbl1 = _make_param_table(with_tbl_header=True)
    hash1 = compute_table_content_hash(tbl1)

    # Change header text — hash should NOT change
    _, tbl2 = _make_param_table(with_tbl_header=True)
    tr0 = tbl2.findall(qn("w:tr"))[0]
    tcs = tr0.findall(qn("w:tc"))
    wt = list(tcs[1].iter(qn("w:t")))[0]
    wt.text = "Point Estimate"

    hash2 = compute_table_content_hash(tbl2)
    assert hash1 == hash2


def test_header_skip_fallback():
    """Fallback header detection: rows before first normal row are excluded."""
    _, tbl1 = _make_param_table(with_tbl_header=False)
    hash1 = compute_table_content_hash(tbl1)

    # Change header text — hash should NOT change
    _, tbl2 = _make_param_table(with_tbl_header=False)
    tr0 = tbl2.findall(qn("w:tr"))[0]
    tcs = tr0.findall(qn("w:tc"))
    wt = list(tcs[1].iter(qn("w:t")))[0]
    wt.text = "Point Estimate"

    hash2 = compute_table_content_hash(tbl2)
    assert hash1 == hash2


def test_full_width_section_label_excluded():
    """Full-width body span (section label) is excluded from hash."""
    _, tbl1 = _make_param_table(with_tbl_header=True)
    hash1 = compute_table_content_hash(tbl1)

    # Change section label — hash should NOT change
    _, tbl2 = _make_param_table(with_tbl_header=True)
    tr1 = tbl2.findall(qn("w:tr"))[1]
    tcs = tr1.findall(qn("w:tc"))
    wt = list(tcs[0].iter(qn("w:t")))[0]
    wt.text = "Fixed Effect Parameters"

    hash2 = compute_table_content_hash(tbl2)
    assert hash1 == hash2


def test_full_width_footer_excluded():
    """Full-width footer (abbreviations) is excluded from hash."""
    _, tbl1 = _make_param_table(with_tbl_header=True)
    hash1 = compute_table_content_hash(tbl1)

    # Change footer text — hash should NOT change
    _, tbl2 = _make_param_table(with_tbl_header=True)
    last_tr = tbl2.findall(qn("w:tr"))[-1]
    tcs = last_tr.findall(qn("w:tc"))
    wt = list(tcs[0].iter(qn("w:t")))[0]
    wt.text = "Abbreviations: CI = confidence intervals; CV = coefficient of variation"

    hash2 = compute_table_content_hash(tbl2)
    assert hash1 == hash2


def test_body_data_change_still_detected():
    """Body data changes are still detected despite header/footer skipping."""
    _, tbl1 = _make_param_table(with_tbl_header=True)
    hash1 = compute_table_content_hash(tbl1)

    # Change a body value — hash SHOULD change
    _, tbl2 = _make_param_table(with_tbl_header=True)
    tr2 = tbl2.findall(qn("w:tr"))[2]
    tcs = tr2.findall(qn("w:tc"))
    wt = list(tcs[3].iter(qn("w:t")))[0]  # "170"
    wt.text = "107"

    hash2 = compute_table_content_hash(tbl2)
    assert hash1 != hash2


def test_header_detection_primary():
    """_detect_header_rows uses w:tblHeader when present."""
    _, tbl = _make_param_table(with_tbl_header=True)
    headers = _detect_header_rows(tbl)
    assert headers == {0}


def test_header_detection_fallback():
    """_detect_header_rows falls back to span-based detection."""
    _, tbl = _make_param_table(with_tbl_header=False)
    headers = _detect_header_rows(tbl)
    # Row 0 has gridSpan, row 1 has gridSpan — first normal row is row 2
    assert 0 in headers
    assert 1 in headers


def test_no_headers_no_skip():
    """Table with no spans has no header rows detected."""
    _, tbl = _make_baseline_table()
    headers = _detect_header_rows(tbl)
    assert headers == set()


def test_logical_column_count():
    """_logical_column_count returns correct count for spanned tables."""
    _, tbl = _make_param_table(with_tbl_header=True)
    assert _logical_column_count(tbl) == 6


def test_full_width_detection():
    """_is_full_width_row correctly identifies full-width rows."""
    _, tbl = _make_param_table(with_tbl_header=True)
    trs = tbl.findall(qn("w:tr"))
    logical_cols = _logical_column_count(tbl)

    # Row 0: header with partial spans — NOT full width
    assert not _is_full_width_row(trs[0], logical_cols)
    # Row 1: section label spanning all 6 cols — full width
    assert _is_full_width_row(trs[1], logical_cols)
    # Row 2: normal data row — NOT full width
    assert not _is_full_width_row(trs[2], logical_cols)
    # Row 4: footer spanning all 6 cols — full width
    assert _is_full_width_row(trs[4], logical_cols)


def test_mixed_table_only_body_hashed():
    """In a mixed table, only body data rows contribute to the hash."""
    _, tbl = _make_param_table(with_tbl_header=True)

    # Build a table with ONLY the body data rows for comparison
    doc2 = Document()
    body_only = doc2.add_table(rows=2, cols=6)
    for r, (texts, _) in enumerate(PARAM_BODY):
        for c, val in enumerate(texts):
            body_only.cell(r, c).paragraphs[0].text = val

    hash_full = compute_table_content_hash(tbl)
    hash_body = compute_table_content_hash(body_only._tbl)

    assert hash_full == hash_body


def test_vmerge_header_detection():
    """vMerge in header rows triggers fallback header detection."""
    doc = Document()
    table = doc.add_table(rows=4, cols=3)
    tbl = table._tbl

    # Row 0: "Dose" with vMerge=restart
    tc0_0 = _get_cell(tbl, 0, 0)
    table.cell(0, 0).paragraphs[0].text = "Dose"
    tcPr0 = tc0_0.find(qn("w:tcPr"))
    if tcPr0 is None:
        tcPr0 = OxmlElement("w:tcPr")
        tc0_0.insert(0, tcPr0)
    vm = OxmlElement("w:vMerge")
    vm.set(qn("w:val"), "restart")
    tcPr0.append(vm)

    table.cell(0, 1).paragraphs[0].text = "Number"
    table.cell(0, 2).paragraphs[0].text = ""

    # Row 1: vMerge continue + sub-headers
    tc1_0 = _get_cell(tbl, 1, 0)
    table.cell(1, 0).paragraphs[0].text = ""
    tcPr1 = tc1_0.find(qn("w:tcPr"))
    if tcPr1 is None:
        tcPr1 = OxmlElement("w:tcPr")
        tc1_0.insert(0, tcPr1)
    vm2 = OxmlElement("w:vMerge")
    tcPr1.append(vm2)

    table.cell(1, 1).paragraphs[0].text = "N"
    table.cell(1, 2).paragraphs[0].text = "OBS"

    # Rows 2-3: normal data
    table.cell(2, 0).paragraphs[0].text = "Placebo"
    table.cell(2, 1).paragraphs[0].text = "10"
    table.cell(2, 2).paragraphs[0].text = "8"
    table.cell(3, 0).paragraphs[0].text = "6 mg"
    table.cell(3, 1).paragraphs[0].text = "12"
    table.cell(3, 2).paragraphs[0].text = "11"

    headers = _detect_header_rows(tbl)
    assert 0 in headers
    assert 1 in headers
    assert 2 not in headers
    assert 3 not in headers


# ── Additional coverage ───────────────────────────────────────────────────


def test_cell_shading_does_not_change_hash():
    """Cell background color / shading should not change hash."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    tc = _get_cell(tbl2, 2, 0)
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = OxmlElement("w:tcPr")
        tc.insert(0, tcPr)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "FFFF00")
    tcPr.append(shd)

    assert compute_table_content_hash(tbl2) == base


def test_multiple_section_labels_excluded():
    """Multiple full-width section labels are all excluded from hash."""
    # Build a table with 3 section labels like a real param table:
    #   Row 0: header with spans
    #   Row 1: "Structural" (full-width)
    #   Row 2: body data
    #   Row 3: "IIV Parameters" (full-width)
    #   Row 4: body data
    #   Row 5: "Residual Error" (full-width)
    #   Row 6: body data
    #   Row 7: footer (full-width)
    doc = Document()
    table = doc.add_table(rows=0, cols=4)
    tbl = table._tbl

    def _add_spanned_row(tbl, text, cols):
        tr = OxmlElement("w:tr")
        tc = OxmlElement("w:tc")
        tcPr = OxmlElement("w:tcPr")
        gs = OxmlElement("w:gridSpan")
        gs.set(qn("w:val"), str(cols))
        tcPr.append(gs)
        tc.append(tcPr)
        p = OxmlElement("w:p")
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        run.append(t)
        p.append(run)
        tc.append(p)
        tr.append(tc)
        tbl.append(tr)

    def _add_header_row(tbl, texts, spans):
        tr = OxmlElement("w:tr")
        trPr = OxmlElement("w:trPr")
        trPr.append(OxmlElement("w:tblHeader"))
        tr.append(trPr)
        for text, span in zip(texts, spans):
            tc = OxmlElement("w:tc")
            if span > 1:
                tcPr = OxmlElement("w:tcPr")
                gs = OxmlElement("w:gridSpan")
                gs.set(qn("w:val"), str(span))
                tcPr.append(gs)
                tc.append(tcPr)
            p = OxmlElement("w:p")
            run = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.text = text
            run.append(t)
            p.append(run)
            tc.append(p)
            tr.append(tc)
        tbl.append(tr)

    def _add_data_row(tbl, texts):
        tr = OxmlElement("w:tr")
        for text in texts:
            tc = OxmlElement("w:tc")
            p = OxmlElement("w:p")
            run = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.text = text
            run.append(t)
            p.append(run)
            tc.append(p)
            tr.append(tc)
        tbl.append(tr)

    _add_header_row(tbl, ["", "Estimate", "CI"], [2, 1, 1])
    _add_spanned_row(tbl, "Structural Parameters", 4)
    _add_data_row(tbl, ["CL", "0.79", "0.75, 0.83", "2.6"])
    _add_spanned_row(tbl, "IIV Parameters", 4)
    _add_data_row(tbl, ["IIV-CL", "0.09", "0.06, 0.12", "16.9"])
    _add_spanned_row(tbl, "Residual Error", 4)
    _add_data_row(tbl, ["Prop", "0.15", "0.12, 0.18", "10.6"])
    _add_spanned_row(tbl, "Abbreviations: CI = confidence intervals", 4)

    hash1 = compute_table_content_hash(tbl)

    # Change all three section labels — hash should NOT change
    trs = tbl.findall(qn("w:tr"))
    for idx in [1, 3, 5]:  # section label rows
        tc = trs[idx].findall(qn("w:tc"))[0]
        wt = list(tc.iter(qn("w:t")))[0]
        wt.text = "CHANGED LABEL"

    # Also change footer
    footer_tc = trs[7].findall(qn("w:tc"))[0]
    wt_footer = list(footer_tc.iter(qn("w:t")))[0]
    wt_footer.text = "CHANGED FOOTER"

    hash2 = compute_table_content_hash(tbl)
    assert hash1 == hash2


def test_simple_table_with_footer_only():
    """Table with no body spans but a full-width footer — footer excluded."""
    doc = Document()
    table = doc.add_table(rows=0, cols=3)
    tbl = table._tbl

    # 2 normal data rows
    for row_texts in [["A", "B", "C"], ["1", "2", "3"]]:
        tr = OxmlElement("w:tr")
        for text in row_texts:
            tc = OxmlElement("w:tc")
            p = OxmlElement("w:p")
            run = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.text = text
            run.append(t)
            p.append(run)
            tc.append(p)
            tr.append(tc)
        tbl.append(tr)

    # Full-width footer
    tr_footer = OxmlElement("w:tr")
    tc_footer = OxmlElement("w:tc")
    tcPr = OxmlElement("w:tcPr")
    gs = OxmlElement("w:gridSpan")
    gs.set(qn("w:val"), "3")
    tcPr.append(gs)
    tc_footer.append(tcPr)
    p = OxmlElement("w:p")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Abbreviations: N = number"
    run.append(t)
    p.append(run)
    tc_footer.append(p)
    tr_footer.append(tc_footer)
    tbl.append(tr_footer)

    hash1 = compute_table_content_hash(tbl)

    # Change footer — hash should NOT change
    wt = list(tc_footer.iter(qn("w:t")))[0]
    wt.text = "Abbreviations: TOTALLY DIFFERENT"

    hash2 = compute_table_content_hash(tbl)
    assert hash1 == hash2

    # Change body data — hash SHOULD change
    body_tc = tbl.findall(qn("w:tr"))[1].findall(qn("w:tc"))[0]
    wt_body = list(body_tc.iter(qn("w:t")))[0]
    wt_body.text = "99"

    hash3 = compute_table_content_hash(tbl)
    assert hash3 != hash2


def test_clear_cell_to_empty():
    """Clearing a populated cell to empty should change hash."""
    _, tbl = _make_baseline_table()
    base = compute_table_content_hash(tbl)

    _, tbl2 = _make_baseline_table()
    tc = _get_cell(tbl2, 3, 2)  # row 4, col 3 = "120"
    wt = _get_wt_elements(tc)[0]
    wt.text = ""

    assert compute_table_content_hash(tbl2) != base


# ── Cell Reconciliation Tests ─────────────────────────────────────────────


class _FakeLogger:
    """Minimal logger for testing reconcile_table_cells."""

    def __init__(self):
        self.messages = []

    def debug(self, msg):
        self.messages.append(("DEBUG", msg))

    def info(self, msg):
        self.messages.append(("INFO", msg))


def testextract_body_grid_matches_hash():
    """extract_body_grid returns data consistent with compute_table_content_hash."""
    _, tbl = _make_param_table(with_tbl_header=True)
    grid = extract_body_grid(tbl)
    canonical = grid_to_canonical(grid)

    import hashlib

    expected_hash = hashlib.sha256(canonical.encode("utf-8")).hexdigest()
    actual_hash = compute_table_content_hash(tbl)
    assert expected_hash == actual_hash


def test_encode_decode_roundtrip():
    """Grid survives gzip+base64 encode/decode round-trip."""
    grid = [
        ["CL/F", "\u03b81", "170", "105, 235"],
        ["Vc/F", "\u03b82", "40", "34, 46"],
    ]
    encoded = encode_body_grid(grid)
    decoded = _decode_body_grid(encoded)
    assert decoded == grid


def test_encode_decode_empty_cells():
    """Encode/decode handles empty cell values."""
    grid = [["A", "", "C"], ["", "B", ""]]
    encoded = encode_body_grid(grid)
    decoded = _decode_body_grid(encoded)
    assert decoded == grid



def test_encode_decode_single_empty_cell_row():
    """A single-cell row with empty string survives roundtrip (C2 regression)."""
    grid = [[""], ["a"]]
    encoded = encode_body_grid(grid)
    decoded = _decode_body_grid(encoded)
    assert decoded == grid


def test_reconcile_same_grid_no_changes():
    """Reconciliation with identical grid updates zero cells."""
    _, tbl = _make_baseline_table()
    grid = extract_body_grid(tbl)
    logger = _FakeLogger()

    result = reconcile_table_cells(tbl, grid, logger)
    assert result is True
    assert any("Reconciled 0 cell(s)" in msg for _, msg in logger.messages)


def test_reconcile_updates_changed_cells():
    """Reconciliation updates only cells that differ."""
    _, tbl = _make_baseline_table()
    original_grid = extract_body_grid(tbl)

    # Mutate one cell in the table
    tc = _get_cell(tbl, 2, 2)  # "100"
    wt = _get_wt_elements(tc)[0]
    wt.text = "999"

    logger = _FakeLogger()
    # Reconcile back to original
    result = reconcile_table_cells(tbl, original_grid, logger)
    assert result is True

    # Verify the cell was restored
    restored_text = "".join(
        t.text for t in tc.iter(qn("w:t")) if t.text
    ).strip()
    assert restored_text == "100"
    assert any("Reconciled 1 cell(s)" in msg for _, msg in logger.messages)


def test_reconcile_preserves_formatting():
    """Reconciliation preserves w:rPr, w:tcPr, w:pPr on updated cells."""
    _, tbl = _make_baseline_table()
    original_grid = extract_body_grid(tbl)

    # Add bold formatting to a cell
    tc = _get_cell(tbl, 2, 0)  # "Revenue"
    runs = _get_run_elements(tc)
    rPr = OxmlElement("w:rPr")
    rPr.append(OxmlElement("w:b"))
    runs[0].insert(0, rPr)

    # Add cell shading
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = OxmlElement("w:tcPr")
        tc.insert(0, tcPr)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "FFFF00")
    tcPr.append(shd)

    # Mutate text
    wt = _get_wt_elements(tc)[0]
    wt.text = "WRONG"

    # Reconcile
    logger = _FakeLogger()
    result = reconcile_table_cells(tbl, original_grid, logger)
    assert result is True

    # Text restored
    restored_text = "".join(
        t.text for t in tc.iter(qn("w:t")) if t.text
    ).strip()
    assert restored_text == "Revenue"

    # Bold still present
    restored_runs = _get_run_elements(tc)
    assert restored_runs[0].find(qn("w:rPr")) is not None
    assert restored_runs[0].find(qn("w:rPr")).find(qn("w:b")) is not None

    # Cell shading still present
    restored_tcPr = tc.find(qn("w:tcPr"))
    assert restored_tcPr is not None
    assert restored_tcPr.find(qn("w:shd")) is not None


def test_reconcile_fails_row_count_mismatch():
    """Reconciliation fails when row counts differ (pathway 1c)."""
    _, tbl = _make_baseline_table()
    short_grid = [["Revenue", "Q1", "100"]]

    logger = _FakeLogger()
    result = reconcile_table_cells(tbl, short_grid, logger)
    assert result is False


def test_reconcile_fails_column_count_mismatch():
    """Reconciliation fails when column counts differ (pathway 1c)."""
    _, tbl = _make_baseline_table()
    grid = extract_body_grid(tbl)
    narrow_grid = [[row[0], row[1]] for row in grid]

    logger = _FakeLogger()
    result = reconcile_table_cells(tbl, narrow_grid, logger)
    assert result is False


def test_reconcile_hash_refreshed():
    """After reconciliation, content hash matches the source grid."""
    _, tbl = _make_baseline_table()
    original_grid = extract_body_grid(tbl)
    original_hash = compute_table_content_hash(tbl)

    # Mutate a cell
    tc = _get_cell(tbl, 2, 2)
    wt = _get_wt_elements(tc)[0]
    wt.text = "999"
    assert compute_table_content_hash(tbl) != original_hash

    # Reconcile
    logger = _FakeLogger()
    reconcile_table_cells(tbl, original_grid, logger)

    # Hash should match original again
    assert compute_table_content_hash(tbl) == original_hash


def test_update_cell_text_preserves_rpr():
    """_update_cell_text preserves run formatting properties."""
    _, tbl = _make_baseline_table()
    tc = _get_cell(tbl, 2, 0)  # "Revenue"

    # Add italic to the run
    runs = _get_run_elements(tc)
    rPr = OxmlElement("w:rPr")
    rPr.append(OxmlElement("w:i"))
    runs[0].insert(0, rPr)

    # Update text
    _update_cell_text(tc, "NewValue")

    # Text changed
    text = "".join(t.text for t in tc.iter(qn("w:t")) if t.text).strip()
    assert text == "NewValue"

    # Italic preserved
    new_runs = _get_run_elements(tc)
    assert len(new_runs) == 1
    assert new_runs[0].find(qn("w:rPr")) is not None
    assert new_runs[0].find(qn("w:rPr")).find(qn("w:i")) is not None


def test_update_cell_text_consolidates_split_runs():
    """_update_cell_text consolidates multiple runs into one."""
    _, tbl = _make_baseline_table()
    tc = _get_cell(tbl, 2, 0)
    p = _get_paragraph_elements(tc)[0]

    # Split into two runs
    for r in list(p.findall(qn("w:r"))):
        p.remove(r)
    for text in ["Rev", "enue"]:
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        run.append(t)
        p.append(run)

    assert len(p.findall(qn("w:r"))) == 2

    _update_cell_text(tc, "Revenue")

    # Should be consolidated to one run
    assert len(p.findall(qn("w:r"))) == 1
    text = "".join(t.text for t in tc.iter(qn("w:t")) if t.text).strip()
    assert text == "Revenue"


def test_reconcile_on_param_table():
    """Reconciliation works on a param-style table with headers and section labels."""
    _, tbl = _make_param_table(with_tbl_header=True)
    original_grid = extract_body_grid(tbl)

    # Mutate a body cell (row 2 = first body row, col 3 = "170")
    trs = tbl.findall(qn("w:tr"))
    body_tr = trs[2]  # first body row
    body_tcs = body_tr.findall(qn("w:tc"))
    wt = list(body_tcs[3].iter(qn("w:t")))[0]
    wt.text = "171"

    logger = _FakeLogger()
    result = reconcile_table_cells(tbl, original_grid, logger)
    assert result is True

    # Verify cell was restored to "170"
    restored_text = "".join(
        t.text for t in body_tcs[3].iter(qn("w:t")) if t.text
    ).strip()
    assert restored_text == "170"
