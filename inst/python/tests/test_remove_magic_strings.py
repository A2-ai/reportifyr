import shutil
import tempfile
from pathlib import Path

from docx import Document

from reportipyr.magic import remove_magic_strings


DATA_DIR = Path(__file__).parent / "data"


def test_remove_magic_strings_preserves_image_and_table():
    source_doc = DATA_DIR / "test-doc-draft.docx"
    docx_in = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    shutil.copyfile(source_doc, docx_in)

    docx_out = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    remove_magic_strings(docx_in, docx_out)

    out_doc = Document(docx_out)

    # Magic string paragraphs should be removed or cleared
    assert all("{rpfy}:" not in p.text for p in out_doc.paragraphs)

    # Image should still exist
    has_image = any(
        run.element.xpath(".//pic:pic")
        for p in out_doc.paragraphs
        for run in p.runs
    )
    assert has_image is True

    # Table should still exist
    assert len(out_doc.tables) >= 1
