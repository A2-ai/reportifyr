import shutil
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from reportipyr.docx_utils import keep_caption_next, remove_bookmarks
from reportipyr.magic import get_magic_pattern
from reportipyr.tables import remove_tables

DATA_DIR = Path(__file__).parent / "data"


def test_remove_tables_removes_table():
    source_doc = DATA_DIR / "test-doc-draft.docx"
    docx_in = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    shutil.copyfile(source_doc, docx_in)

    # Confirm input has at least one table
    assert len(Document(docx_in).tables) >= 1

    docx_out = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    remove_tables(docx_in, docx_out)

    out_doc = Document(docx_out)
    assert len(out_doc.tables) == 0


def test_remove_bookmarks_removes_fp_bookmarks():
    source_doc = DATA_DIR / "test-doc-draft.docx"
    docx_in = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    shutil.copyfile(source_doc, docx_in)

    # Confirm input has fp_ bookmarks
    in_doc = Document(docx_in)
    fp_bookmarks = [
        el for el in in_doc.element.findall(
            ".//w:bookmarkStart", namespaces=in_doc.element.nsmap
        )
        if (el.get(qn("w:name")) or "").startswith("fp_")
    ]
    assert len(fp_bookmarks) >= 1

    docx_out = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    remove_bookmarks(docx_in, docx_out)

    out_doc = Document(docx_out)
    remaining = [
        el for el in out_doc.element.findall(
            ".//w:bookmarkStart", namespaces=out_doc.element.nsmap
        )
        if (el.get(qn("w:name")) or "").startswith("fp_")
    ]
    assert len(remaining) == 0


def test_keep_caption_next_sets_property():
    source_doc = DATA_DIR / "test-doc-draft.docx"
    docx_in = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    shutil.copyfile(source_doc, docx_in)

    docx_out = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    keep_caption_next(docx_in, docx_out)

    out_doc = Document(docx_out)
    magic_pattern = get_magic_pattern()

    # Caption paragraph(s) should have keepNext
    caption_has_keep_next = any(
        p._element.xpath(".//w:keepNext")
        for p in out_doc.paragraphs
        if (p.style and p.style.name == "Caption")
        or any(
            "SEQ Table" in instr.text or "SEQ Figure" in instr.text
            for instr in p._element.xpath(".//w:instrText")
        )
    )
    assert caption_has_keep_next is True

    # Magic string paragraph(s) following a caption should also have keepNext
    magic_has_keep_next = any(
        p._element.xpath(".//w:keepNext")
        for p in out_doc.paragraphs
        if magic_pattern.search(p.text)
    )
    assert magic_has_keep_next is True
