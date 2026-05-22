import tempfile
from pathlib import Path

from docx import Document

from reportipyr.tables_xml import add_table_xml, W_TBL, W_P


SIMPLE_TBL = """<?xml version="1.0" encoding="UTF-8"?>
<w:tbl xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:tblPr/>
  <w:tblGrid>
    <w:gridCol w:w="2000"/>
    <w:gridCol w:w="2000"/>
  </w:tblGrid>
  <w:tr>
    <w:tc><w:p><w:r><w:t>A</w:t></w:r></w:p></w:tc>
    <w:tc><w:p><w:r><w:t>B</w:t></w:r></w:p></w:tc>
  </w:tr>
  <w:tr>
    <w:tc><w:p><w:r><w:t>1</w:t></w:r></w:p></w:tc>
    <w:tc><w:p><w:r><w:t>2</w:t></w:r></w:p></w:tc>
  </w:tr>
</w:tbl>
"""


def _write_xml(table_dir: Path, name: str, content: str = SIMPLE_TBL) -> Path:
    path = table_dir / name
    path.write_text(content)
    return path


def _make_docx_with_magic(magic_text: str) -> str:
    docx_in = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    doc = Document()
    doc.add_paragraph(magic_text)
    doc.save(docx_in)
    return docx_in


def _out_docx() -> str:
    return tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name


def _count_tbls(doc: Document) -> int:
    return len(doc.element.body.xpath(".//w:tbl"))


def test_add_table_xml_inserts_tbl_after_magic_paragraph():
    table_dir = Path(tempfile.mkdtemp())
    _write_xml(table_dir, "demographics.xml")

    docx_in = _make_docx_with_magic("{rpfy}:demographics.xml")
    docx_out = _out_docx()

    add_table_xml(docx_in, docx_out, str(table_dir))

    out_doc = Document(docx_out)
    assert _count_tbls(out_doc) == 1

    # The first body-level <w:tbl> should immediately follow the magic paragraph.
    body = out_doc.element.body
    children = list(body)
    tbl_idx = next(i for i, c in enumerate(children) if c.tag == W_TBL)
    assert children[tbl_idx - 1].tag.endswith("}p")
    assert "{rpfy}:demographics.xml" in "".join(
        t.text or "" for t in children[tbl_idx - 1].xpath(".//w:t")
    )


def test_add_table_xml_skips_when_file_missing():
    table_dir = Path(tempfile.mkdtemp())  # empty
    docx_in = _make_docx_with_magic("{rpfy}:absent.xml")
    docx_out = _out_docx()

    add_table_xml(docx_in, docx_out, str(table_dir))

    out_doc = Document(docx_out)
    assert _count_tbls(out_doc) == 0


def test_add_table_xml_skips_when_tbl_already_present():
    table_dir = Path(tempfile.mkdtemp())
    _write_xml(table_dir, "demographics.xml")

    docx_in = _make_docx_with_magic("{rpfy}:demographics.xml")
    docx_out = _out_docx()

    add_table_xml(docx_in, docx_out, str(table_dir))
    # Second pass: should be idempotent, no second <w:tbl> added.
    docx_out2 = _out_docx()
    add_table_xml(docx_out, docx_out2, str(table_dir))

    out_doc = Document(docx_out2)
    assert _count_tbls(out_doc) == 1


def test_add_table_xml_rejects_non_tbl_root():
    table_dir = Path(tempfile.mkdtemp())
    bad = """<?xml version="1.0"?>
<root xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:p/>
</root>
"""
    _write_xml(table_dir, "bad.xml", bad)

    docx_in = _make_docx_with_magic("{rpfy}:bad.xml")
    docx_out = _out_docx()

    add_table_xml(docx_in, docx_out, str(table_dir))

    out_doc = Document(docx_out)
    assert _count_tbls(out_doc) == 0


def test_add_table_xml_rejects_malformed_xml():
    table_dir = Path(tempfile.mkdtemp())
    (table_dir / "broken.xml").write_text("<w:tbl><not-closed>")

    docx_in = _make_docx_with_magic("{rpfy}:broken.xml")
    docx_out = _out_docx()

    add_table_xml(docx_in, docx_out, str(table_dir))

    out_doc = Document(docx_out)
    assert _count_tbls(out_doc) == 0


def test_add_table_xml_ignores_non_xml_extensions():
    table_dir = Path(tempfile.mkdtemp())
    (table_dir / "demographics.csv").write_text("a,b\n1,2\n")

    docx_in = _make_docx_with_magic("{rpfy}:demographics.csv")
    docx_out = _out_docx()

    add_table_xml(docx_in, docx_out, str(table_dir))

    out_doc = Document(docx_out)
    assert _count_tbls(out_doc) == 0


def test_add_table_xml_inserts_into_cell_with_trailing_paragraph():
    table_dir = Path(tempfile.mkdtemp())
    _write_xml(table_dir, "nested.xml")

    docx_in = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = "{rpfy}:nested.xml"
    doc.save(docx_in)

    docx_out = _out_docx()
    add_table_xml(docx_in, docx_out, str(table_dir))

    out_doc = Document(docx_out)
    # The original table has 1 nested table inside its single cell.
    tbls = out_doc.element.body.xpath(".//w:tbl")
    assert len(tbls) == 2  # outer + inserted

    # Inserted <w:tbl> must live inside a <w:tc>, and the cell must still end
    # with <w:p> (OOXML requirement).
    inserted = next(
        t for t in tbls
        if t.getparent() is not None and t.getparent().tag.endswith("}tc")
    )
    tc = inserted.getparent()
    assert tc[-1].tag == W_P
