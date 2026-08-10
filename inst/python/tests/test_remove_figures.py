import json
import os
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image

from reportipyr.figures import remove_figures


def _make_png() -> str:
    path = tempfile.NamedTemporaryFile(suffix=".png", delete=False).name
    img = Image.new("RGB", (10, 10), color=(255, 0, 0))
    img.save(path)
    return path


def _write_yaml(contents: str) -> str:
    path = tempfile.NamedTemporaryFile(suffix=".yaml", delete=False).name
    Path(path).write_text(contents)
    return path


# ── skip_unchanged fixtures ────────────────────────────────────────────────

FIGURE_HASH = "a1b2c3d4e5f6"


def _write_figure_metadata(
    figure_dir: Path, figure: str = "figure.png", hash_value: str = FIGURE_HASH
) -> None:
    """Write the sibling <name>_<ext>_metadata.json load_artifact_hash reads."""
    object_name, extension = os.path.splitext(figure)
    metadata_file = figure_dir / f"{object_name}_{extension[1:]}_metadata.json"
    metadata_file.write_text(json.dumps({"object_meta": {"hash": hash_value}}))


def _set_drawing_alt_text(paragraph, alt_text: str) -> None:
    """Set descr on every wp:docPr in the paragraph's inline drawings."""
    for doc_pr in paragraph._element.xpath(".//wp:docPr"):
        doc_pr.set("descr", alt_text)


def _make_unchanged_docx(
    magic_text: str,
    width: float,
    height: float,
    in_cell: bool = False,
) -> tuple[str, str]:
    """Build a docx whose drawing carries a [hash:...] matching its metadata.

    Returns (docx_in, figure_dir) — with skip_unchanged: true this is the
    "artifact unchanged, leave the figure alone" shape.

    The magic-string run is bolded so a test can tell a no-op run from a
    destructive ``paragraph.text = ...`` rewrite that happens to produce the
    same characters.
    """
    img_path = _make_png()

    figure_dir = Path(tempfile.mkdtemp())
    _write_figure_metadata(figure_dir)

    doc = Document()
    if in_cell:
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        magic_par = cell.paragraphs[0]
        magic_par.text = magic_text
        fig_par = cell.add_paragraph()
    else:
        magic_par = doc.add_paragraph(magic_text)
        fig_par = doc.add_paragraph()

    magic_par.runs[0].bold = True
    fig_par.add_run().add_picture(
        img_path, width=Inches(width), height=Inches(height)
    )
    _set_drawing_alt_text(fig_par, f"A figure [hash:{FIGURE_HASH}]")

    docx_in = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    doc.save(docx_in)
    return docx_in, str(figure_dir)


def _has_drawing(doc: Document) -> bool:
    return any(p._element.xpath(".//w:drawing") for p in doc.paragraphs)


def _cell_paragraphs(doc: Document):
    return doc.tables[0].cell(0, 0).paragraphs


def test_remove_figures_updates_magic_string_with_dimensions():
    img_path = _make_png()

    docx_in = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    doc = Document()
    doc.add_paragraph("{rpfy}:figure.png")
    fig_par = doc.add_paragraph()
    fig_par.add_run().add_picture(img_path, width=Inches(2), height=Inches(3))
    doc.save(docx_in)

    config_yaml = _write_yaml("use_embedded_dimensions: true\n")

    docx_out = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    remove_figures(docx_in, docx_out, config_yaml)

    out_doc = Document(docx_out)
    text = out_doc.paragraphs[0].text
    assert "width: 2.0" in text
    assert "height: 3.0" in text


def test_remove_figures_no_dimensions_when_disabled():
    img_path = _make_png()

    docx_in = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    doc = Document()
    doc.add_paragraph("{rpfy}:figure.png")
    fig_par = doc.add_paragraph()
    fig_par.add_run().add_picture(img_path, width=Inches(2), height=Inches(3))
    doc.save(docx_in)

    config_yaml = _write_yaml("use_embedded_dimensions: false\n")

    docx_out = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    remove_figures(docx_in, docx_out, config_yaml)

    out_doc = Document(docx_out)
    # Magic string should be unchanged
    assert out_doc.paragraphs[0].text.strip() == "{rpfy}:figure.png"
    # Image paragraph should be removed
    has_drawing = any(
        p._element.xpath(".//w:drawing") for p in out_doc.paragraphs
    )
    assert has_drawing is False


def test_remove_figures_multi_figure_reconstructs_magic_string():
    img_path = _make_png()

    docx_in = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    doc = Document()
    doc.add_paragraph("{rpfy}:[a.png, b.png]")
    doc.add_paragraph().add_run().add_picture(img_path, width=Inches(2), height=Inches(3))
    doc.add_paragraph().add_run().add_picture(img_path, width=Inches(4), height=Inches(5))
    doc.save(docx_in)

    config_yaml = _write_yaml("use_embedded_dimensions: true\n")

    docx_out = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    remove_figures(docx_in, docx_out, config_yaml)

    out_doc = Document(docx_out)
    text = out_doc.paragraphs[0].text

    # Should have bracket syntax with both figures and their dimensions
    assert text.startswith("{rpfy}:[")
    assert text.endswith("]")
    assert "a.png<width: 2.0, height: 3.0>" in text
    assert "b.png<width: 4.0, height: 5.0>" in text

    # Images should be removed
    has_drawing = any(
        p._element.xpath(".//w:drawing") for p in out_doc.paragraphs
    )
    assert has_drawing is False


def test_skip_unchanged_syncs_resized_dimensions():
    """A user resize is synced back into the magic string even when the
    artifact hash is unchanged and the figure is retained."""
    docx_in, figure_dir = _make_unchanged_docx(
        "{rpfy}:figure.png<width: 2.0, height: 3.0>", width=4, height=5
    )

    config_yaml = _write_yaml(
        "skip_unchanged: true\nuse_embedded_dimensions: true\n"
    )

    docx_out = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    remove_figures(docx_in, docx_out, config_yaml, figure_dir=figure_dir)

    out_doc = Document(docx_out)
    text = out_doc.paragraphs[0].text
    assert "width: 4.0" in text
    assert "height: 5.0" in text
    assert "2.0" not in text
    assert "3.0" not in text

    # Regression guard: the retained figure must not be deleted
    assert _has_drawing(out_doc) is True


def test_skip_unchanged_no_rewrite_when_size_matches():
    """When the drawing already matches the magic string, the paragraph is
    left completely alone (no destructive rewrite of its runs)."""
    magic_text = "{rpfy}:figure.png<width: 2.0, height: 3.0>"
    docx_in, figure_dir = _make_unchanged_docx(magic_text, width=2, height=3)

    config_yaml = _write_yaml(
        "skip_unchanged: true\nuse_embedded_dimensions: true\n"
    )

    docx_out = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    remove_figures(docx_in, docx_out, config_yaml, figure_dir=figure_dir)

    out_doc = Document(docx_out)
    magic_par = out_doc.paragraphs[0]
    assert magic_par.text == magic_text

    # Run-level formatting survives, so the paragraph was never rewritten
    assert len(magic_par.runs) == 1
    assert magic_par.runs[0].bold is True

    assert _has_drawing(out_doc) is True


def test_skip_unchanged_syncs_resized_dimensions_in_cell():
    """Same sync, for a magic string and drawing inside a table cell."""
    docx_in, figure_dir = _make_unchanged_docx(
        "{rpfy}:figure.png<width: 2.0, height: 3.0>",
        width=4,
        height=5,
        in_cell=True,
    )

    config_yaml = _write_yaml(
        "skip_unchanged: true\nuse_embedded_dimensions: true\n"
    )

    docx_out = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    remove_figures(docx_in, docx_out, config_yaml, figure_dir=figure_dir)

    out_doc = Document(docx_out)
    cell_paras = _cell_paragraphs(out_doc)
    text = cell_paras[0].text
    assert "width: 4.0" in text
    assert "height: 5.0" in text
    assert "2.0" not in text
    assert "3.0" not in text

    # Regression guard: the retained cell figure must not be deleted
    has_drawing = any(
        p._element.xpath(".//w:drawing") for p in cell_paras
    )
    assert has_drawing is True


def test_skip_unchanged_respects_use_embedded_dimensions_false():
    """use_embedded_dimensions: false still leaves the magic string alone."""
    magic_text = "{rpfy}:figure.png<width: 2.0, height: 3.0>"
    docx_in, figure_dir = _make_unchanged_docx(magic_text, width=4, height=5)

    config_yaml = _write_yaml(
        "skip_unchanged: true\nuse_embedded_dimensions: false\n"
    )

    docx_out = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    remove_figures(docx_in, docx_out, config_yaml, figure_dir=figure_dir)

    out_doc = Document(docx_out)
    assert out_doc.paragraphs[0].text == magic_text
    assert _has_drawing(out_doc) is True
